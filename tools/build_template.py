"""Dựng file mẫu nhập liệu chuẩn cho app Bảng chỉ số TBH hàng không.

Chạy:  python3 tools/build_template.py templates/Mau_nhap_lieu_UY_x_update_ddmmyyyy.xlsx
Tên sheet và tiêu đề cột giữ đúng như file theo dõi hiện tại để app đọc được.
"""
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Mau_nhap_lieu_UY_x_update_ddmmyyyy.xlsx'
N = 300                      # số dòng dựng sẵn công thức / kiểm tra
LAST = N + 1
FN = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864'); HF = Font(name=FN, bold=True, color='FFFFFF', size=10)
HDR_F = PatternFill('solid', fgColor='595959')            # tiêu đề cột công thức
INP = Font(name=FN, color='0000FF', size=10)              # ô nhập: chữ xanh dương
FRM = Font(name=FN, color='000000', size=10)              # ô công thức: chữ đen
GRAY = PatternFill('solid', fgColor='F2F2F2')             # nền ô công thức
REQ = PatternFill('solid', fgColor='FFF2CC')              # thiếu dữ liệu bắt buộc
RED = PatternFill('solid', fgColor='F8CBAD'); ORG = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='D9D9D9'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
DT = 'dd/mm/yyyy'; USD = '#,##0.00'; PCT = '0.00'
EX = 'VÍ DỤ – xoá dòng này'

wb = Workbook()

# ---------------- Danh mục ----------------
LISTS = {
    'L_Status': ['Bound', 'Offered', 'Cancelled', 'Declined'],
    'L_Broker': ['AJG', 'Marsh', 'WTW', 'McGill', 'MBPH', 'KMD', 'Lockton', 'GC', 'UIB', 'Prudent', 'Howden', 'Aon', 'BMS', 'Gallagher Re'],
    'L_Region': ['Asia Pacific RIM', 'Africa', 'Middle East', 'Europe', 'Americas'],
    'L_Channel': ['A1', 'MAAPS', 'Open market'],
    'L_Leader': ['AXA XL', 'Allianz', 'Arrow', 'Atrium', 'Beazley', 'Chubb', 'Convex', 'Global', 'Liberty', 'LRA', 'Nexus', 'SFMIC', 'SIEL', 'Starr', 'Tokio Marine Kiln'],
    'L_RiskType': ['HSL', 'LIABS', 'HW', 'TLO'],
    'L_Type': ['RI', 'XOL', 'Direct', 'Fronting'],
    'L_Category': ['ABC', '2nd Schedule'],
    'L_Flying': ['Flying', 'Non-Flying'],
    'L_Source': ['Major PVI', 'Non-major Alif'],
    'L_PremStatus': ['Estimated', 'Closing', 'Adjusted'],
    'Hull': ['Engine', 'Ground Handling', 'Hard landing', 'Bird strike', 'Spare', 'Partial loss', 'Total loss', 'Other'],
    'Liability': ['Passenger', 'Human', 'Third party property', 'Ground Handling', 'Product / Hangarkeepers', 'Other'],
    'L_Coverage': ['Hull', 'Liability'],
    'L_ClaimStatus': ['Open', 'Closed', 'Reopened'],
    'L_PaidStatus': ['Paid', 'Partial', 'Unpaid'],
    'L_InstPaid': ['Paid', 'Unpaid'],
}
wl = wb.active; wl.title = 'Danh mục'
for j, (name, vals) in enumerate(LISTS.items(), 1):
    c = wl.cell(row=1, column=j, value=name); c.font = HF; c.fill = HDR
    wl.column_dimensions[L(j)].width = max(14, len(name) + 4)
    for i, v in enumerate(vals, 2):
        wl.cell(row=i, column=j, value=v).font = INP
    # vùng tên để dài thêm 30 dòng cho phép bổ sung giá trị mới
    ref = f"'Danh mục'!${L(j)}$2:${L(j)}${len(vals) + 31}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
wl['A35'] = 'Bổ sung giá trị mới vào cuối cột tương ứng (ô trống ngay dưới danh sách). Danh sách chọn trong các sheet tự cập nhật.'
wl['A35'].font = Font(name=FN, italic=True, size=9)

def header(ws, cols, formula_cols=(), widths=None, notes=None):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = HF; c.fill = HDR_F if j in formula_cols else HDR
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BOX
        if notes and h in notes: c.comment = Comment(notes[h], 'Mẫu nhập liệu')
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{L(len(cols))}{LAST}'
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[L(j)].width = (widths or {}).get(j, 14)

def style_rows(ws, ncols, formula_cols, fmt=None):
    for r in range(2, LAST + 1):
        for j in range(1, ncols + 1):
            c = ws.cell(row=r, column=j)
            c.font = FRM if j in formula_cols else INP
            if j in formula_cols: c.fill = GRAY
            if fmt and j in fmt: c.number_format = fmt[j]

def dv_list(ws, name, rng, prompt=None, stop=True):
    d = DataValidation(type='list', formula1=f'={name}', allow_blank=True, showErrorMessage=True,
                       errorStyle='stop' if stop else 'warning', errorTitle='Giá trị không có trong danh mục',
                       error='Chọn từ danh sách. Muốn thêm giá trị mới, bổ sung vào sheet "Danh mục".')
    if prompt: d.promptTitle = 'Hướng dẫn'; d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(rng)

def dv_custom(ws, formula, rng, title, msg, stop=True, prompt=None):
    d = DataValidation(type='custom', formula1=formula, allow_blank=True, showErrorMessage=True,
                       errorStyle='stop' if stop else 'warning', errorTitle=title, error=msg)
    if prompt: d.promptTitle = 'Hướng dẫn'; d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(rng)

def dv_date(ws, rng, prompt='Nhập ngày dạng dd/mm/yyyy (ô ngày của Excel), không gõ chữ.'):
    d = DataValidation(type='date', operator='between', formula1='DATE(2015,1,1)', formula2='DATE(2040,12,31)', allow_blank=True,
                       showErrorMessage=True, errorTitle='Ngày không hợp lệ', error='Nhập ngày thật trong khoảng 2015–2040, dạng dd/mm/yyyy. Chưa biết thì để trống, không điền ngày tạm.')
    d.promptTitle = 'Ngày'; d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(rng)

def dv_num(ws, rng, lo, hi, msg, whole=False):
    d = DataValidation(type='whole' if whole else 'decimal', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                       showErrorMessage=True, errorTitle='Số không hợp lệ', error=msg)
    ws.add_data_validation(d); d.add(rng)

def col_rng(c): return f'{c}2:{c}{LAST}'

# ---------------- Policy ----------------
wp = wb.create_sheet('Policy', 0)
pcols = ['MRO', 'Broker', 'PIAS Ref', 'PVI ref number', 'Original Insured', 'Cedant Name', 'Region', 'Domicile', 'Risk offer date', 'Status',
         'Risk bound date', 'Inception', 'Expiry', 'Days on risk', 'Days remain', 'Flying/Non-flying', 'Leader', 'Risk type', 'Type', 'Category',
         'Channel', 'Prior PVI ref', 'Endorsement of', 'Source', 'Notes']
pf = {14, 15}
notes = {
    'MRO': 'Mã MRO của broker. Bắt buộc.',
    'PVI ref number': 'Số tham chiếu PVI – DUY NHẤT cho mỗi dòng. Bắt buộc.',
    'Inception': 'Ngày hiệu lực. Bắt buộc.',
    'Expiry': 'Ngày hết hạn – phải sau ngày hiệu lực. Bắt buộc.',
    'Risk bound date': 'Ngày bind – không được sau ngày hiệu lực, trừ khi có xác nhận không tổn thất (NKLL).',
    'Prior PVI ref': 'Số tham chiếu PVI của hợp đồng năm trước (nếu là tái tục). Dùng để tính tỷ lệ tái tục chính xác.',
    'Endorsement of': 'Chỉ điền cho dòng sửa đổi bổ sung (AP/RP): số tham chiếu PVI của hợp đồng gốc.',
    'Source': 'Major PVI (PVI tự khai thác) hoặc Non-major Alif.',
    'Notes': 'Ghi chú nội bộ – không ghi vào các cột số tham chiếu.',
    'Days on risk': 'Công thức – không nhập.', 'Days remain': 'Công thức – không nhập.',
}
header(wp, pcols, pf, {1: 20, 3: 18, 4: 16, 5: 36, 6: 30, 7: 16, 8: 14, 22: 16, 23: 16, 24: 15, 25: 30}, notes)
style_rows(wp, len(pcols), pf, {9: DT, 11: DT, 12: DT, 13: DT, 14: '0', 15: '0'})
for r in range(2, LAST + 1):
    wp[f'N{r}'] = f'=IF(OR(L{r}="",M{r}=""),"",MAX(0,MIN(TODAY(),M{r})-L{r}+1))'
    wp[f'O{r}'] = f'=IF(M{r}="","",M{r}-TODAY())'
dv_list(wp, 'L_Broker', col_rng('B'), 'Chọn broker')
dv_custom(wp, 'COUNTIF($D$2:$D$%d,D2)=1' % LAST, col_rng('D'), 'Số tham chiếu bị trùng', 'PVI ref number phải duy nhất. Sửa đổi bổ sung dùng số riêng và ghi số gốc vào cột "Endorsement of".')
dv_list(wp, 'L_Region', col_rng('G'), 'Chọn khu vực')
dv_date(wp, col_rng('I'))
dv_list(wp, 'L_Status', col_rng('J'), 'Bound / Offered / Cancelled / Declined')
dv_custom(wp, 'OR(K2="",L2="",K2<=L2)', col_rng('K'), 'Bind sau ngày hiệu lực', 'Ngày bind sau ngày hiệu lực: chỉ chấp nhận khi có xác nhận không tổn thất (NKLL). Bấm Yes để giữ, ghi lý do vào cột Notes.', stop=False, prompt='Nhập ngày bind dạng dd/mm/yyyy; không được sau ngày hiệu lực.')
dv_date(wp, col_rng('L'))
dv_custom(wp, 'AND(ISNUMBER(M2),M2>L2,M2<=L2+550)', col_rng('M'), 'Ngày hết hạn không hợp lệ', 'Ngày hết hạn phải sau ngày hiệu lực và không quá 18 tháng. Kiểm tra lại năm.', prompt='Ngày hết hạn dd/mm/yyyy, sau ngày hiệu lực.')
dv_list(wp, 'L_Flying', col_rng('P'))
dv_list(wp, 'L_Leader', col_rng('Q'), 'Chọn leader')
dv_list(wp, 'L_RiskType', col_rng('R')); dv_list(wp, 'L_Type', col_rng('S')); dv_list(wp, 'L_Category', col_rng('T'))
dv_list(wp, 'L_Channel', col_rng('U'), 'A1 / MAAPS / Open market')
dv_list(wp, 'L_Source', col_rng('X'))
rng = f'A2:{L(len(pcols))}{LAST}'
wp.conditional_formatting.add(f'M2:M{LAST}', FormulaRule(formula=['AND(M2<>"",L2<>"",M2<=L2)'], fill=RED))
wp.conditional_formatting.add(f'K2:K{LAST}', FormulaRule(formula=['AND(K2<>"",L2<>"",K2>L2)'], fill=ORG))
for c in ['D', 'E', 'J', 'L', 'M', 'B', 'U']:
    wp.conditional_formatting.add(f'{c}2:{c}{LAST}', FormulaRule(formula=[f'AND($A2<>"",{c}2="")'], fill=REQ))

# ---------------- Risk ----------------
wr = wb.create_sheet('Risk', 1)
rcols = ['MRO', 'Broker', 'Original Insured', 'PVI share', 'Capacity deployment', 'Hull', 'Spare limit', 'CSL', 'PA', 'AVN52E', 'A/C', 'AFV', 'Pax', 'Dept', 'Spare value']
rf = {1, 2, 3, 5}
header(wr, rcols, rf, {1: 20, 3: 36, 5: 18, 12: 16}, {'PVI share': 'Tỷ lệ PVI tính theo %: 0,5 nghĩa là 0,5%. Hợp đồng Bound phải > 0.',
                                                      'MRO': 'Tự lấy từ sheet Policy cùng dòng – không nhập.', 'Capacity deployment': 'Công thức: share × (Hull + CSL).',
                                                      'A/C': 'Số tàu (số nguyên).', 'AFV': 'Tổng giá trị đội tàu (USD).'})
style_rows(wr, len(rcols), rf, {4: '0.000', 5: '#,##0', 6: '#,##0', 7: '#,##0', 8: '#,##0', 9: '#,##0', 10: '#,##0', 11: '0', 12: '#,##0', 13: '#,##0', 14: '#,##0', 15: '#,##0'})
for r in range(2, LAST + 1):
    wr[f'A{r}'] = f'=IF(Policy!A{r}="","",Policy!A{r})'
    wr[f'B{r}'] = f'=IF(Policy!A{r}="","",Policy!B{r})'
    wr[f'C{r}'] = f'=IF(Policy!A{r}="","",Policy!E{r})'
    wr[f'E{r}'] = f'=IF(A{r}="","",D{r}*(F{r}+H{r})/100)'
dv_num(wr, col_rng('D'), 0, 100, 'Tỷ lệ PVI từ 0 đến 100 (%). Ví dụ 0,5 nghĩa là 0,5%.')
for c in 'FGHIJLMNO': dv_num(wr, col_rng(c), 0, 1e12, 'Nhập số tiền USD ≥ 0.')
dv_num(wr, col_rng('K'), 0, 3000, 'Số tàu là số nguyên từ 0 đến 3.000. Kiểm tra không đảo cột với AFV.', whole=True)
wr.conditional_formatting.add(f'D2:D{LAST}', FormulaRule(formula=['AND(Policy!$J2="Bound",OR(D2="",D2=0))'], fill=RED))
wr.conditional_formatting.add(f'L2:L{LAST}', FormulaRule(formula=['AND(K2>0,L2/K2>500000000)'], fill=ORG))

# ---------------- Premium ----------------
wm = wb.create_sheet('Premium', 2)
mcols = ['MRO', 'Broker', 'Original Insured', 'Region', 'Domicile', 'PVI share', 'Deposit GWP', 'Deposit NWP', 'PVI deposite GWP', 'PVI deposite NWP',
         'Brokerage', 'SMB', 'Total deduction', 'Facility fee', 'Other Deduction', 'Fund', 'PVI actual earning', 'Diffirental from lead', 'Facility',
         'Month', 'Year', 'Policy status', 'Premium status']
mf = {1, 2, 3, 4, 5, 6, 8, 9, 10, 13, 17, 19, 20, 21, 22}
header(wm, mcols, mf, {1: 20, 3: 36, 7: 15, 8: 15, 9: 15, 10: 15, 17: 16}, {
    'Deposit GWP': 'Phí gốc 100% (USD). Bắt buộc với hợp đồng Bound.', 'Brokerage': '% brokerage trên GWP, ví dụ 21,5.', 'SMB': '% SMB trên GWP.',
    'Facility fee': '% facility fee trên NWP.', 'Other Deduction': 'Khấu trừ khác (USD).', 'Fund': 'Khoản quỹ (USD).',
    'Diffirental from lead': '% chênh lệch so với leader (nếu có).',
    'PVI deposite GWP': 'Công thức: Deposit GWP × share.', 'PVI deposite NWP': 'Công thức: PVI GWP × (1 − tổng khấu trừ %).',
    'PVI actual earning': 'Công thức: NWP × (1 − facility fee %) − khấu trừ khác − quỹ.', 'Total deduction': 'Công thức: brokerage + SMB.'})
style_rows(wm, len(mcols), mf, {6: '0.000', 7: USD, 8: USD, 9: USD, 10: USD, 11: PCT, 12: PCT, 13: PCT, 14: PCT, 15: USD, 16: USD, 17: USD, 18: PCT})
for r in range(2, LAST + 1):
    g = f'IF(Policy!A{r}="",""'
    wm[f'A{r}'] = f'={g},Policy!A{r})'
    wm[f'B{r}'] = f'={g},Policy!B{r})'
    wm[f'C{r}'] = f'={g},Policy!E{r})'
    wm[f'D{r}'] = f'={g},Policy!G{r})'
    wm[f'E{r}'] = f'={g},Policy!H{r})'
    wm[f'F{r}'] = f'={g},Risk!D{r})'
    wm[f'H{r}'] = f'=IF(A{r}="","",G{r}*(1-M{r}/100))'
    wm[f'I{r}'] = f'=IF(A{r}="","",G{r}*F{r}/100)'
    wm[f'J{r}'] = f'=IF(A{r}="","",I{r}*(1-M{r}/100))'
    wm[f'M{r}'] = f'=IF(A{r}="","",K{r}+L{r})'
    wm[f'Q{r}'] = f'=IF(A{r}="","",J{r}*(1-N{r}/100)-O{r}-P{r})'
    wm[f'S{r}'] = f'={g},Policy!U{r})'
    wm[f'T{r}'] = f'=IF(OR(A{r}="",Policy!L{r}=""),"",MONTH(Policy!L{r}))'
    wm[f'U{r}'] = f'=IF(OR(A{r}="",Policy!L{r}=""),"",YEAR(Policy!L{r}))'
    wm[f'V{r}'] = f'={g},Policy!J{r})'
dv_num(wm, col_rng('G'), -1e10, 1e10, 'Nhập phí gốc 100% (USD). Hoàn phí nhập số âm.')
for c in 'KLN': dv_num(wm, col_rng(c), 0, 60, 'Nhập phần trăm từ 0 đến 60, ví dụ 21,5.')
for c in 'OP': dv_num(wm, col_rng(c), 0, 1e9, 'Nhập số tiền USD ≥ 0.')
dv_list(wm, 'L_PremStatus', col_rng('W'))
wm.conditional_formatting.add(f'G2:G{LAST}', FormulaRule(formula=['AND($V2="Bound",OR(G2="",G2=0))'], fill=REQ))

# ---------------- Claim ----------------
wc = wb.create_sheet('Claim', 3)
ccols = ['MRO', 'Broker', 'Original Insured', 'Claim number', 'Date claim reported', 'Occurrence date', 'Location', 'Loss details', 'Coverage', 'Type of loss',
         'Loss adjuster/ Lawyer', 'Paid Claim (100%)', 'Paid Fee (100%)', 'Reserved claim', 'Resered Fee', 'Gross Incurred', 'PVI share', 'PVI Paid Claim',
         'PVI Paid Fee', 'PVI reserved claim', 'PVI reserved Fee', "PVI's Gross Incured", 'Status', 'Lastest Movement Date', 'Refund amout', 'Paid Status']
cf = {2, 3, 16, 17, 18, 19, 20, 21, 22}
header(wc, ccols, cf, {1: 20, 3: 32, 4: 20, 8: 40, 22: 16}, {
    'MRO': 'Chọn MRO của hợp đồng đang hiệu lực vào ngày tổn thất (đúng UY).', 'Claim number': 'Số hồ sơ bồi thường – bắt buộc, duy nhất.',
    'Occurrence date': 'Ngày tổn thất dd/mm/yyyy. Tổn thất kéo dài: ghi ngày đầu, mô tả khoảng thời gian ở Loss details.',
    'Coverage': 'Hull hoặc Liability.', 'Type of loss': 'Danh sách đổi theo Coverage.',
    'Paid Claim (100%)': 'Số đã trả 100% (USD).', 'Reserved claim': 'Dự phòng 100% (USD). Hồ sơ đang mở phải có dự phòng ban đầu.',
    'PVI share': 'Tự lấy từ sheet Risk theo MRO.', "PVI's Gross Incured": 'Công thức: tổng phần PVI đã trả + dự phòng.',
    'Paid Status': 'Paid chỉ khi PVI đã thực trả.'})
style_rows(wc, len(ccols), cf, {5: DT, 6: DT, 12: USD, 13: USD, 14: USD, 15: USD, 16: USD, 17: '0.000', 18: USD, 19: USD, 20: USD, 21: USD, 22: USD, 24: DT, 25: USD})
for r in range(2, LAST + 1):
    m = f'MATCH($A{r},Policy!$A$2:$A${LAST},0)'
    wc[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(Policy!$B$2:$B${LAST},{m}),"MRO không có"))'
    wc[f'C{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(Policy!$E$2:$E${LAST},{m}),"MRO không có"))'
    wc[f'P{r}'] = f'=IF($A{r}="","",L{r}+M{r}+N{r}+O{r})'
    wc[f'Q{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(Risk!$D$2:$D${LAST},{m}),0))'
    wc[f'R{r}'] = f'=IF($A{r}="","",L{r}*Q{r}/100)'
    wc[f'S{r}'] = f'=IF($A{r}="","",M{r}*Q{r}/100)'
    wc[f'T{r}'] = f'=IF($A{r}="","",N{r}*Q{r}/100)'
    wc[f'U{r}'] = f'=IF($A{r}="","",O{r}*Q{r}/100)'
    wc[f'V{r}'] = f'=IF($A{r}="","",R{r}+S{r}+T{r}+U{r})'
d = DataValidation(type='list', formula1=f'=Policy!$A$2:$A${LAST}', allow_blank=True, showErrorMessage=True, errorTitle='MRO không có trong Policy', error='Chọn MRO đã nhập ở sheet Policy.')
wc.add_data_validation(d); d.add(col_rng('A'))
dv_custom(wc, 'COUNTIF($D$2:$D$%d,D2)=1' % LAST, col_rng('D'), 'Số hồ sơ bị trùng', 'Mỗi hồ sơ bồi thường có một số riêng.')
dv_date(wc, col_rng('E')); dv_date(wc, col_rng('F')); dv_date(wc, col_rng('X'))
dv_list(wc, 'L_Coverage', col_rng('I'))
d = DataValidation(type='list', formula1='=INDIRECT($I2)', allow_blank=True, showErrorMessage=True, errorTitle='Loại tổn thất không phù hợp', error='Chọn Coverage trước; loại tổn thất phải thuộc danh sách của Coverage đó.')
wc.add_data_validation(d); d.add(col_rng('J'))
for c in 'LMNOY': dv_num(wc, col_rng(c), -1e10, 1e10, 'Nhập số tiền USD.')
dv_list(wc, 'L_ClaimStatus', col_rng('W')); dv_list(wc, 'L_PaidStatus', col_rng('Z'))
inc = f'INDEX(Policy!$L$2:$L${LAST},MATCH($A2,Policy!$A$2:$A${LAST},0))'
exp = f'INDEX(Policy!$M$2:$M${LAST},MATCH($A2,Policy!$A$2:$A${LAST},0))'
wc.conditional_formatting.add(f'F2:F{LAST}', FormulaRule(formula=[f'AND(F2<>"",IFERROR(OR(F2<{inc},F2>{exp}),FALSE))'], fill=RED))
wc.conditional_formatting.add(f'E2:E{LAST}', FormulaRule(formula=['AND(E2<>"",F2<>"",E2<F2)'], fill=ORG))
wc.conditional_formatting.add(f'N2:N{LAST}', FormulaRule(formula=['AND($A2<>"",$W2="Open",$P2=0)'], fill=REQ))
wc.conditional_formatting.add(f'Z2:Z{LAST}', FormulaRule(formula=['AND(Z2="Paid",R2+S2=0)'], fill=ORG))
for c in ['D', 'F', 'I', 'W']:
    wc.conditional_formatting.add(f'{c}2:{c}{LAST}', FormulaRule(formula=[f'AND($A2<>"",{c}2="")'], fill=REQ))

# ---------------- Installment ----------------
wi = wb.create_sheet('Installment', 4)
icols = ['MRO', 'Broker', 'Original Insured', 'Closing Ref', 'Closing Date']
for n in range(1, 13): icols += [f'Inst {n} Date', f'Inst {n} Share %', f'Inst {n} Amount', f'Inst {n} Due Date', f'Inst {n} Paid', f'Inst {n} Actual NWP']
icols.append('Last Movement Date')
ifc = {2, 3}
header(wi, icols, ifc, {1: 20, 3: 32}, {'MRO': 'Chọn MRO đã nhập ở Policy. Mọi hợp đồng Bound đều phải có lịch, kể cả trả một lần.',
                                        'Inst 1 Amount': 'Số tiền kỳ (USD, theo NWP).', 'Inst 1 Due Date': 'Hạn thanh toán dd/mm/yyyy.',
                                        'Inst 1 Paid': 'Chọn Paid khi đã thu, đồng thời nhập số thực thu.', 'Inst 1 Actual NWP': 'Số thực thu (USD).'})
fmt = {5: DT, len(icols): DT}
for n in range(12):
    b = 6 + n * 6
    fmt.update({b: DT, b + 1: '0.00', b + 2: USD, b + 3: DT, b + 5: USD})
style_rows(wi, len(icols), ifc, fmt)
for r in range(2, LAST + 1):
    m = f'MATCH($A{r},Policy!$A$2:$A${LAST},0)'
    wi[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(Policy!$B$2:$B${LAST},{m}),"MRO không có"))'
    wi[f'C{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(Policy!$E$2:$E${LAST},{m}),"MRO không có"))'
d = DataValidation(type='list', formula1=f'=Policy!$A$2:$A${LAST}', allow_blank=True, showErrorMessage=True, errorTitle='MRO không có trong Policy', error='Chọn MRO đã nhập ở sheet Policy.')
wi.add_data_validation(d); d.add(col_rng('A'))
dv_date(wi, col_rng('E')); dv_date(wi, col_rng(L(len(icols))))
for n in range(12):
    b = 6 + n * 6
    dc, ac, duc, pc, nc = L(b), L(b + 2), L(b + 3), L(b + 4), L(b + 5)
    dv_date(wi, col_rng(dc)); dv_date(wi, col_rng(duc))
    dv_list(wi, 'L_InstPaid', col_rng(pc))
    dv_num(wi, col_rng(ac), -1e9, 1e9, 'Nhập số tiền kỳ (USD).'); dv_num(wi, col_rng(nc), -1e9, 1e9, 'Nhập số thực thu (USD).')
    wi.conditional_formatting.add(col_rng(duc), FormulaRule(formula=[f'AND({duc}2<>"",{duc}2<TODAY(),{pc}2<>"Paid")'], fill=RED))
    wi.conditional_formatting.add(col_rng(duc), FormulaRule(formula=[f'AND({duc}2<>"",{duc}2>=TODAY(),{duc}2<=TODAY()+30,{pc}2<>"Paid")'], fill=ORG))
    wi.conditional_formatting.add(col_rng(nc), FormulaRule(formula=[f'AND({pc}2="Paid",OR({nc}2="",{nc}2=0))'], fill=REQ))

# ---------------- Dòng ví dụ (Status = Offered nên app không tính vào số liệu) ----------------
d0 = datetime.datetime
ex_p = ['B0000EXAMPLE01', 'AJG', '', 'PVI202600001', EX, 'Example Insurance Co', 'Asia Pacific RIM', 'Vietnam', d0(2026, 9, 1), 'Offered', d0(2026, 9, 20), d0(2026, 10, 1), d0(2027, 9, 30),
        None, None, 'Flying', 'AXA XL', 'HSL', 'RI', 'ABC', 'A1', 'PVI202500001', '', 'Major PVI', 'Dòng ví dụ – Status Offered nên app không tính']
for j, v in enumerate(ex_p, 1):
    if v is not None and j not in pf: wp.cell(row=2, column=j, value=v)
for c, v in {'D': 0.5, 'F': 100000000, 'G': 0, 'H': 1000000000, 'K': 12, 'L': 400000000}.items(): wr[f'{c}2'] = v
for c, v in {'G': 2000000, 'K': 21.5, 'L': 3.5, 'N': 0, 'O': 0, 'P': 0, 'W': 'Estimated'}.items(): wm[f'{c}2'] = v
for c, v in {'A': 'B0000EXAMPLE01', 'D': 'PVIAVC2026000001', 'E': d0(2026, 11, 5), 'F': d0(2026, 11, 1), 'H': 'Mô tả tổn thất', 'I': 'Hull', 'J': 'Engine',
             'L': 0, 'M': 0, 'N': 500000, 'O': 0, 'W': 'Open', 'X': d0(2026, 11, 5)}.items(): wc[f'{c}2'] = v
for c, v in {'A': 'B0000EXAMPLE01', 'F': d0(2026, 10, 1), 'H': 3750, 'I': d0(2026, 11, 1), 'J': 'Paid', 'K': 3750, 'L': d0(2027, 4, 1), 'N': 3750, 'O': d0(2027, 5, 1), 'P': 'Unpaid'}.items(): wi[f'{c}2'] = v

# ---------------- Hướng dẫn ----------------
wh = wb.create_sheet('Hướng dẫn', 0)
wh.column_dimensions['A'].width = 3; wh.column_dimensions['B'].width = 26; wh.column_dimensions['C'].width = 100
wh['B2'] = 'Mẫu nhập liệu sổ theo dõi nhận tái bảo hiểm hàng không'; wh['B2'].font = Font(name=FN, bold=True, size=14)
wh['B3'] = 'Dùng cho app Bảng chỉ số TBH hàng không – https://dungexim-spec.github.io/aviation-ri-kpi/'; wh['B3'].font = Font(name=FN, italic=True, size=10, color='595959')
rows = [
    ('MÀU Ô', ''),
    ('Chữ xanh dương', 'Ô nhập liệu.'),
    ('Nền xám, chữ đen', 'Ô công thức – không nhập, không xoá. Tiêu đề cột công thức có nền xám đậm.'),
    ('Nền vàng', 'Thiếu dữ liệu bắt buộc.'),
    ('Nền đỏ', 'Lỗi cần sửa: ngày hết hạn trước ngày hiệu lực, tổn thất ngoài kỳ hợp đồng, hợp đồng Bound có tỷ lệ PVI = 0, kỳ phí quá hạn chưa thu.'),
    ('Nền cam', 'Cần kiểm tra: bind sau ngày hiệu lực, ngày thông báo trước ngày tổn thất, kỳ phí đến hạn trong 30 ngày, "Paid" nhưng chưa có số tiền.'),
    ('', ''),
    ('QUY TẮC CHÍNH', ''),
    ('1. Một hợp đồng một dòng', 'Mỗi dòng Policy có một PVI ref number duy nhất. Sửa đổi bổ sung (AP/RP) dùng số riêng và ghi số hợp đồng gốc vào cột "Endorsement of".'),
    ('2. Thứ tự dòng', 'Chỉ nhập MRO / broker / tên ở sheet Policy. Sheet Risk và Premium tự lấy theo cùng dòng – dòng 5 của Risk và Premium là của hợp đồng ở dòng 5 Policy. Không chèn, xoá hay sắp xếp dòng riêng ở Risk hoặc Premium.'),
    ('3. Ngày tháng', 'Nhập dạng ngày của Excel (dd/mm/yyyy). Không gõ chữ, không điền ngày tạm 01/01/2025; chưa biết thì để trống.'),
    ('4. Danh sách chọn', 'Broker, Region, Status, Channel, Leader, Coverage, Type of loss, Paid… chọn từ danh sách. Thêm giá trị mới ở sheet "Danh mục".'),
    ('5. Tỷ lệ PVI', 'Nhập theo %: 0,5 nghĩa là 0,5%. Hợp đồng Bound phải > 0.'),
    ('6. Phí', 'Nhập Deposit GWP (100%), Brokerage %, SMB %, Facility fee %, khấu trừ khác, quỹ. GWP, NWP và net earning của PVI tự tính.'),
    ('7. Bồi thường', 'Chọn MRO của hợp đồng đang hiệu lực vào ngày tổn thất. Nhập số 100%; phần PVI tự tính. Mỗi hồ sơ có số riêng; hồ sơ mở phải có dự phòng.'),
    ('8. Thu phí', 'Mọi hợp đồng Bound đều có lịch ở sheet Installment, kể cả trả một lần. Khi thu: chọn Paid và nhập số thực thu.'),
    ('9. Tái tục', 'Hợp đồng tái tục ghi số PVI ref năm trước vào cột "Prior PVI ref".'),
    ('10. Ghi chú', 'Ghi chú nội bộ để ở cột Notes, không viết vào cột số tham chiếu.'),
    ('', ''),
    ('ĐẶT TÊN FILE', 'UY_<số>_update_<ddmmyyyy>.xlsx, ví dụ UY_2_update_01112026.xlsx (UY2, số liệu chốt ngày 01/11/2026).'),
    ('DÒNG VÍ DỤ', 'Dòng 2 của mỗi sheet là ví dụ (Status = Offered nên app không tính). Có thể ghi đè hoặc xoá.'),
    ('SỨC CHỨA', f'Dựng sẵn công thức và kiểm tra cho {N} hợp đồng. Cần thêm: chọn dòng cuối, kéo xuống (fill down) ở tất cả các sheet.'),
]
r = 5
for a, b in rows:
    ca = wh.cell(row=r, column=2, value=a); cb = wh.cell(row=r, column=3, value=b)
    ca.font = Font(name=FN, bold=bool(a) and (b == '' or a.isupper()), size=10); cb.font = Font(name=FN, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical='top'); ca.alignment = Alignment(vertical='top', wrap_text=True)
    if a and b == '': ca.fill = PatternFill('solid', fgColor='D9E1F2'); wh.cell(row=r, column=3).fill = PatternFill('solid', fgColor='D9E1F2')
    r += 1
samples = [(6, None, None), (7, GRAY, FRM), (8, REQ, None), (9, RED, None), (10, ORG, None)]
for row, fill, font in samples:
    c = wh.cell(row=row, column=2)
    if fill: c.fill = fill
    if row == 6: c.font = Font(name=FN, color='0000FF', bold=True, size=10)

wb.active = 0
wb.save(OUT)
print('saved', OUT)
